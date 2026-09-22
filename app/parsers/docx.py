import io

from docx import Document

from app.parsers.pdf import Block, table_to_text

_HEADING_MARKERS = ("heading", "заголовок", "title", "название")


def _is_heading(style_name: str) -> bool:
    name = (style_name or "").lower()
    return any(marker in name for marker in _HEADING_MARKERS)


def parse_docx(data: bytes) -> list[Block]:
    document = Document(io.BytesIO(data))
    blocks: list[Block] = []
    section = ""

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if _is_heading(paragraph.style.name):
            section = text
        blocks.append(Block("text", text, page=1, section=section))

    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        blocks.append(
            Block("table", table_to_text(rows), page=1, section=section, rows=rows)
        )

    return blocks
