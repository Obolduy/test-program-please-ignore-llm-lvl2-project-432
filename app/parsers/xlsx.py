import io

from openpyxl import load_workbook

from app.parsers.pdf import Block

ARTICUL_COLUMNS = ("артикул", "код", "sku")
SECTION = "Спецификация"


def parse_xlsx(data: bytes, sheet: str | None = None) -> list[Block]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook.active
    rows = [
        [("" if cell is None else str(cell)).strip() for cell in row]
        for row in worksheet.iter_rows(values_only=True)
    ]
    workbook.close()

    header_at = next((i for i, row in enumerate(rows) if any(row)), None)
    if header_at is None:
        return []
    header = rows[header_at]

    blocks: list[Block] = []
    for row in rows[header_at + 1 :]:
        if not any(row):
            continue
        pairs = {name: value for name, value in zip(header, row) if name and value}
        body = "; ".join(f"{name}: {value}" for name, value in pairs.items())
        articul = next(
            (value for name, value in pairs.items() if name.lower() in ARTICUL_COLUMNS), ""
        )
        blocks.append(
            Block(
                kind="table",
                content=f"{articul}: {body}" if articul else body,
                page=1,
                section=SECTION,
                rows=[header, row],
            )
        )
    return blocks
